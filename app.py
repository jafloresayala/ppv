import streamlit as st
import pandas as pd
import os
import glob
import re
import plotly.express as px
from io import BytesIO
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt
from datetime import datetime
from ai_chat import render_chat_panel

# ─── Configuración ────────────────────────────────────────────────────────────
FOLDER_PATH   = r"C:\Users\K90016277\PyCharmMiscProject\Backup"
OUTPUT_FILE   = r"C:\Users\K90016277\PyCharmMiscProject\Components_Report_Consolidated.parquet"
METADATA_FILE = r"C:\Users\K90016277\PyCharmMiscProject\processed_files.parquet"
TARGET_SHEET  = "All Customers"

st.set_page_config(page_title="Price Roll Dashboard", layout="wide", initial_sidebar_state="expanded")

if "consolidated_df" not in st.session_state:
    st.session_state["consolidated_df"] = None
if "chat_open" not in st.session_state:
    st.session_state["chat_open"] = True

st.markdown("""
<style>
    .stApp { background: linear-gradient(180deg, #f5f7fb 0%, #eef3f8 100%); }
    .block-container { padding-top: 1.5rem; padding-bottom: 2rem; max-width: 1450px; }
    div[data-testid="stMetric"] {
        background: rgba(255,255,255,0.92);
        border: 1px solid rgba(35,52,77,0.08);
        border-radius: 16px;
        padding: 0.75rem 0.9rem;
        box-shadow: 0 8px 24px rgba(36,61,89,0.06);
    }
    div[data-testid="stDataFrame"] { border-radius: 16px; overflow: hidden; }
</style>
""", unsafe_allow_html=True)


# ─── Capa de datos ────────────────────────────────────────────────────────────

def get_excel_files_robust():
    files = set()
    try:
        for entry in os.scandir(FOLDER_PATH):
            if entry.is_file(follow_symlinks=False):
                name = entry.name
                if not name.startswith("~$") and name.lower().endswith((".xlsx", ".xls", ".xlsm")):
                    files.add(name)
    except (OSError, PermissionError):
        for path in (glob.glob(os.path.join(FOLDER_PATH, "*.xlsx")) +
                     glob.glob(os.path.join(FOLDER_PATH, "*.xls")) +
                     glob.glob(os.path.join(FOLDER_PATH, "*.xlsm"))):
            name = os.path.basename(path)
            if not name.startswith("~$"):
                files.add(name)
    return sorted(files)


def load_metadata():
    if os.path.exists(METADATA_FILE):
        df = pd.read_parquet(METADATA_FILE)
        for col in ["file_name", "status", "reason", "processed_at"]:
            if col not in df.columns:
                df[col] = None
        return df[["file_name", "status", "reason", "processed_at"]]
    return pd.DataFrame(columns=["file_name", "status", "reason", "processed_at"])


def save_metadata(df_meta):
    df_meta.to_parquet(METADATA_FILE, index=False)
    st.cache_data.clear()


def upsert_metadata(df_meta, file_name, status, reason=""):
    df_meta = df_meta[df_meta["file_name"] != file_name]
    return pd.concat([df_meta, pd.DataFrame([{
        "file_name": file_name, "status": status, "reason": reason,
        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }])], ignore_index=True)


# ─── Lógica de negocio ────────────────────────────────────────────────────────

def process_files():
    df_meta = load_metadata()
    processed_ok = set(df_meta.loc[df_meta["status"] == "PROCESADO_OK", "file_name"].dropna())
    new_dfs, ok_count, err_count = [], 0, 0

    for file_name in get_excel_files_robust():
        if file_name in processed_ok:
            continue
        file_path = os.path.join(FOLDER_PATH, file_name)
        match = re.search(r'_(\d{8})', file_name)
        if not match:
            df_meta = upsert_metadata(df_meta, file_name, "ERROR_SIN_FECHA", "No se encontró fecha en el nombre")
            st.warning(f"⚠ {file_name}: sin fecha en el nombre")
            err_count += 1
            continue
        report_date = datetime.strptime(match.group(1), "%Y%m%d").date()
        try:
            xl = pd.ExcelFile(file_path)
            if TARGET_SHEET not in xl.sheet_names:
                df_meta = upsert_metadata(df_meta, file_name, "ERROR_HOJA_FALTANTE", f"Sin hoja '{TARGET_SHEET}'")
                st.warning(f"⚠ {file_name}: sin hoja '{TARGET_SHEET}'")
                err_count += 1
                continue
            df = pd.read_excel(file_path, sheet_name=TARGET_SHEET)
            df.insert(0, "Report_Date", report_date)
            new_dfs.append(df)
            df_meta = upsert_metadata(df_meta, file_name, "PROCESADO_OK")
            st.success(f"✅ {file_name}")
            ok_count += 1
        except Exception as e:
            df_meta = upsert_metadata(df_meta, file_name, "ERROR_PROCESAMIENTO", str(e))
            st.error(f"❌ {file_name}: {e}")
            err_count += 1

    save_metadata(df_meta)
    if not new_dfs:
        return None, ok_count, err_count
    return pd.concat(new_dfs, ignore_index=True, sort=False), ok_count, err_count


def validate_pending_files():
    df_meta = load_metadata()
    latest = (df_meta.sort_values("processed_at").drop_duplicates("file_name", keep="last")
              if not df_meta.empty else pd.DataFrame())
    records = []
    for file_name in get_excel_files_robust():
        row = latest[latest["file_name"] == file_name] if not latest.empty else pd.DataFrame()
        if row.empty:
            records.append({"file_name": file_name, "status": "PENDIENTE_NUNCA_PROCESADO", "reason": "Nunca procesado"})
        elif row.iloc[0]["status"] != "PROCESADO_OK":
            records.append({"file_name": file_name, "status": row.iloc[0]["status"], "reason": row.iloc[0]["reason"]})
    return pd.DataFrame(records)


def get_diagnostics_report():
    df_meta = load_metadata()
    latest = (df_meta.sort_values("processed_at").drop_duplicates("file_name", keep="last")
              if not df_meta.empty else pd.DataFrame())
    rows = []
    for file_name in get_excel_files_robust():
        file_path = os.path.join(FOLDER_PATH, file_name)
        exists = os.path.exists(file_path)
        size_mb = round(os.path.getsize(file_path) / 1_048_576, 2) if exists else 0
        row = latest[latest["file_name"] == file_name] if not latest.empty else pd.DataFrame()
        status = row.iloc[0]["status"] if not row.empty else "SIN_PROCESAR"
        rows.append({"Archivo": file_name, "En Carpeta": "✅" if exists else "❌", "MB": size_mb, "Estado": status})
    return pd.DataFrame(rows)


def clean_dataframe(df):
    """Aplica todos los filtros de limpieza. Retorna (df_clean, metrics_dict)."""
    n0 = len(df)
    df = df[df["Material Type"] != "ROH"]
    n1 = len(df)
    df = df[df["Demand Comments"] != "Not found"]
    n2 = len(df)
    if "PO price comment" in df.columns:
        df = df[df["PO price comment"] != "PO price not found"]
    if "PO price comments" in df.columns:
        df = df[df["PO price comments"] != "PO price not found"]
    n3 = len(df)
    pattern = r'\b(?:OBS|IN)\w*'
    df = df[~df["Component Description"].str.contains(pattern, regex=True, na=False)].copy().reset_index(drop=True)
    n4 = len(df)
    return df, {
        "rows_initial":    n0,
        "removed_material": n0 - n1,
        "removed_demand":   n1 - n2,
        "removed_po":       n2 - n3,
        "removed_obs":      n3 - n4,
    }


# ─── Exportación PowerPoint ───────────────────────────────────────────────────

def build_dashboard_export_data(df, mode):
    required = {"Report_Date", "Component", "Standard Price", "Info Record Price"}
    missing = required.difference(df.columns)
    if missing:
        return {"error": ", ".join(sorted(missing))}

    df = df.copy()
    for col in ["Standard Price", "Info Record Price"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["Report_Date"] = pd.to_datetime(df["Report_Date"], errors="coerce")
    df = df.dropna(subset=list(required))
    if df.empty:
        return {"error": "No hay datos válidos para exportar."}

    total = df["Component"].nunique()
    has_desc = "Component Description" in df.columns
    if has_desc:
        df["Component Description"] = df["Component Description"].fillna("")

    status_df = (
        df.assign(d=df["Standard Price"] - df["Info Record Price"])
        .groupby("Component", as_index=False)
        .agg(
            Has_Gain=("d", lambda v: (v > 0).any()),
            Has_Loss=("d", lambda v: (v < 0).any()),
            Is_Neutral=("d", lambda v: (v == 0).all()),
        )
    )

    if mode == "loss":
        df["Amount"] = df["Info Record Price"] - df["Standard Price"]
        df["Percent"] = (df["Amount"] / df["Standard Price"]) * 100
        df = df[df["Amount"] > 0].copy()
        count   = int(status_df["Has_Loss"].sum())
        neutral = int(status_df["Is_Neutral"].sum())
        metrics = {
            "Componentes totales": total, "Con perdida": count, "Neutrales": neutral,
            "Total perdidas":  df["Amount"].sum() if not df.empty else 0,
            "Perdida promedio": df["Amount"].mean() if not df.empty else 0,
            "Perdida max":     df["Amount"].max() if not df.empty else 0,
        }
        title = "Dashboard de perdida por componente"
        acol, pcol, mcol, tlabel = "Average_Loss", "Average_Loss_Percent", "Max_Loss", "perdida"
    else:
        df["Amount"] = df["Standard Price"] - df["Info Record Price"]
        df["Percent"] = (df["Amount"] / df["Info Record Price"]) * 100
        df = df[df["Amount"] > 0].copy()
        count   = int(status_df["Has_Gain"].sum())
        neutral = int(status_df["Is_Neutral"].sum())
        metrics = {
            "Componentes totales": total, "Con ganancia": count, "Neutrales": neutral,
            "Total ganancias":   df["Amount"].sum() if not df.empty else 0,
            "Ganancia promedio": df["Amount"].mean() if not df.empty else 0,
            "Ganancia max":      df["Amount"].max() if not df.empty else 0,
        }
        title = "Dashboard de ganancia por componente"
        acol, pcol, mcol, tlabel = "Average_Gain", "Average_Gain_Percent", "Max_Gain", "ganancia"

    if df.empty:
        return {"error": "No hay datos para esta vista."}

    df = df.sort_values(["Component", "Report_Date"])
    df["Standard Price Previous"] = df.groupby("Component")["Standard Price"].shift(1)
    df["Variacion Porcentual"] = ((df["Standard Price"] - df["Standard Price Previous"]) / df["Standard Price Previous"]) * 100

    aggs = {
        acol: ("Amount", "mean"), pcol: ("Percent", "mean"), mcol: ("Amount", "max"),
        "Average_Standard_Price":   ("Standard Price", "mean"),
        "Average_Info_Record_Price": ("Info Record Price", "mean"),
        "Average_Variacion_Porcentual": ("Variacion Porcentual", "mean"),
        "Records": ("Component", "size"),
    }
    if has_desc:
        aggs["Component Description"] = ("Component Description", lambda v: next((x for x in v if x), ""))

    summary = (df.groupby("Component", as_index=False).agg(**aggs)
               .sort_values(acol, ascending=False).reset_index(drop=True))

    non_event = max(total - count - neutral, 0)
    return {
        "title": title, "metrics": metrics, "summary_df": summary.head(5),
        "amount_col": acol, "percent_col": pcol, "max_col": mcol,
        "table_title": f"Top componentes con mayor {tlabel} promedio",
        "pie_df": pd.DataFrame({
            "Categoria": [f"Componentes con {tlabel}", "Componentes neutrales", f"Componentes sin {tlabel}"],
            "Cantidad":  [count, neutral, non_event],
        }),
    }


def build_pie_chart_image(pie_df, title, color_map):
    fig, ax = plt.subplots(figsize=(5.0, 3.6), dpi=170)
    colors = [color_map.get(l, "#9fb3c8") for l in pie_df["Categoria"]]
    ax.pie(pie_df["Cantidad"], labels=pie_df["Categoria"],
           autopct=lambda p: f"{p:.1f}%" if p > 0 else "", startangle=90,
           colors=colors, wedgeprops={"width": 0.48, "edgecolor": "white"},
           textprops={"fontsize": 9, "color": "#18324d"})
    ax.set_title(title, fontsize=13, fontweight="bold", color="#18324d", pad=14)
    ax.axis("equal")
    fig.patch.set_facecolor("white")
    buf = BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def build_bar_chart_image(summary_df, amount_col, title, color):
    plot_df = summary_df.iloc[::-1].copy()
    fig, ax = plt.subplots(figsize=(7.0, 3.8), dpi=170)
    bars = ax.barh(plot_df["Component"].astype(str), plot_df[amount_col], color=color, alpha=0.92)
    ax.set_title(title, fontsize=13, fontweight="bold", color="#18324d", pad=14)
    ax.set_xlabel("Monto", fontsize=10, color="#18324d")
    for spine in ["top", "right"]:
        ax.spines[spine].set_visible(False)
    ax.spines["left"].set_color("#d4dce6")
    ax.spines["bottom"].set_color("#d4dce6")
    ax.grid(axis="x", linestyle="--", alpha=0.22, color="#5d6d81")
    ax.set_axisbelow(True)
    ax.tick_params(axis="both", labelsize=9, colors="#18324d")
    for bar in bars:
        ax.text(bar.get_width(), bar.get_y() + bar.get_height() / 2,
                f" {bar.get_width():,.2f}", va="center", fontsize=8, color="#18324d")
    fig.patch.set_facecolor("white")
    buf = BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png", bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def add_dashboard_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9.0), Inches(0.6))
    tf = tb.text_frame
    tf.text = data["title"]
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.bold = True

    xs, ys = [0.5, 2.1, 3.7], [1.0, 2.25]
    fill, line = RGBColor(245, 247, 251), RGBColor(214, 223, 232)
    for i, (label, value) in enumerate(list(data["metrics"].items())[:6]):
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                                      Inches(xs[i % 3]), Inches(ys[i // 3]), Inches(1.45), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = fill
        card.line.color.rgb = line
        tf = card.text_frame
        tf.clear()
        p1 = tf.paragraphs[0]
        p1.text = label
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(93, 109, 129)
        p2 = tf.add_paragraph()
        p2.text = f"{value:,.2f}" if isinstance(value, float) else f"{value:,}"
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(24, 50, 77)

    if data["summary_df"].empty:
        return

    is_loss = "perdida" in data["title"].lower()
    key = "perdida" if is_loss else "ganancia"
    pie_colors = {
        f"Componentes con {key}":  "#d62728" if is_loss else "#2ca02c",
        "Componentes neutrales":   "#9fb3c8",
        f"Componentes sin {key}":  "#2ca02c" if is_loss else "#d62728",
    }
    pie_img = build_pie_chart_image(data["pie_df"], f"Participacion — {key}", pie_colors)
    bar_img = build_bar_chart_image(data["summary_df"], data["amount_col"],
                                    data["table_title"] + " Top 5", "#d62728" if is_loss else "#2ca02c")
    slide.shapes.add_picture(pie_img, Inches(3.8), Inches(1.1), Inches(3.9), Inches(2.9))
    slide.shapes.add_picture(bar_img, Inches(7.6), Inches(1.05), Inches(5.3), Inches(3.05))


def build_powerpoint_export(df):
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for mode in ("loss", "gain"):
        data = build_dashboard_export_data(df, mode)
        if "error" in data:
            label = "perdida" if mode == "loss" else "ganancia"
            add_dashboard_slide(prs, {
                "title": f"Dashboard de {label} por componente",
                "metrics": {"Error": data["error"]}, "summary_df": pd.DataFrame(),
                "amount_col": "Amount", "percent_col": "Percent", "max_col": "Max",
                "table_title": data["error"], "pie_df": pd.DataFrame(),
            })
        else:
            add_dashboard_slide(prs, data)
    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


# ─── Secciones de UI ──────────────────────────────────────────────────────────

def render_overview_panel():
    files    = get_excel_files_robust()
    df_meta  = load_metadata()
    ok, errors, pending = 0, 0, 0
    if not df_meta.empty:
        latest  = df_meta.sort_values("processed_at").drop_duplicates("file_name", keep="last")
        ok      = int((latest["status"] == "PROCESADO_OK").sum())
        errors  = int((latest["status"] != "PROCESADO_OK").sum())
        pending = len([f for f in files if f not in set(latest["file_name"])])
    else:
        pending = len(files)
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Archivos en carpeta",    len(files))
    c2.metric("Procesados",             ok)
    c3.metric("Pendientes / Error",     errors + pending)
    c4.metric("Consolidado", "Disponible" if os.path.exists(OUTPUT_FILE) else "No existe")


def render_control_de_carga():
    st.subheader("Control de carga")
    st.caption("Consolida los archivos .xlsx de la carpeta y valida pendientes antes de continuar.")
    render_overview_panel()

    col1, col2 = st.columns(2)
    if col1.button("Ejecutar consolidación", use_container_width=True):
        new_data, ok, err = process_files()
        if new_data is None:
            st.info("No hay archivos nuevos para consolidar.")
        else:
            if os.path.exists(OUTPUT_FILE):
                new_data = pd.concat([pd.read_parquet(OUTPUT_FILE), new_data], ignore_index=True)
            new_data.to_parquet(OUTPUT_FILE, index=False)
            st.success(f"Consolidación completada — {ok} procesados, {err} con error.")

    if col2.button("Validar pendientes", use_container_width=True):
        pending_df = validate_pending_files()
        if pending_df.empty:
            st.success("Al día. No hay archivos pendientes.")
        else:
            st.warning(f"{len(pending_df)} archivo(s) pendiente(s) o con error.")
            st.dataframe(pending_df, use_container_width=True, hide_index=True)

    st.markdown("---")
    st.subheader("Diagnóstico de archivos")
    st.caption("Verificación de cada archivo detectado vs su estado en metadata.")
    diag_df = get_diagnostics_report()
    if diag_df.empty:
        st.info("No hay archivos detectados.")
    else:
        c1, c2 = st.columns([1, 4])
        c1.metric("Detectados", len(diag_df))
        c1.metric("En carpeta", len(diag_df[diag_df["En Carpeta"] == "✅"]))
        c2.dataframe(diag_df, use_container_width=True, hide_index=True)
        missing = diag_df[diag_df["En Carpeta"] == "❌"]
        if not missing.empty:
            st.error(f"{len(missing)} archivo(s) detectado(s) pero no encontrado(s) en carpeta.")
            st.dataframe(missing[["Archivo"]], use_container_width=True, hide_index=True)


def render_consolidado():
    st.subheader("Consolidado")
    st.caption("Carga el parquet consolidado para revisar la base antes del análisis.")
    if not os.path.exists(OUTPUT_FILE):
        st.info("Aún no existe un archivo consolidado.")
        return
    if st.button("Leer consolidado", use_container_width=True):
        st.session_state["consolidated_df"] = pd.read_parquet(OUTPUT_FILE)
    df = st.session_state["consolidated_df"]
    if df is not None:
        c1, c2, c3 = st.columns(3)
        c1.metric("Filas",        f"{len(df):,}")
        c2.metric("Columnas",     f"{len(df.columns):,}")
        c3.metric("Componentes",  f"{df['Component'].nunique():,}" if "Component" in df.columns else "N/D")
        with st.expander("Vista previa", expanded=False):
            st.dataframe(df, use_container_width=True)


def render_limpieza(df_clean, df_before_obs, metrics):
    st.subheader("Limpieza de datos")
    st.caption("Eliminamos filas inválidas: Material Type ROH, sin Demanda, sin PO Price y componentes obsoletos.")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Filas iniciales",            f"{metrics['rows_initial']:,}")
    c2.metric("Eliminadas por Material/Demand", f"{metrics['removed_material'] + metrics['removed_demand']:,}")
    c3.metric("Eliminadas por PO not found", f"{metrics['removed_po']:,}")
    c4.metric("Eliminadas por OBS/IN",       f"{metrics['removed_obs']:,}")

    with st.expander("Vista previa después de filtros base", expanded=False):
        st.dataframe(df_before_obs, use_container_width=True)

    st.caption(f"**Filas finales: {len(df_clean):,}**")
    with st.expander("Vista previa dataset listo para dashboard", expanded=False):
        st.dataframe(df_clean, use_container_width=True)


def render_loss_dashboard(df):
    required = {"Report_Date", "Component", "Standard Price", "Info Record Price", "PO Price"}
    missing  = required.difference(df.columns)
    if missing:
        st.warning("Faltan columnas: " + ", ".join(sorted(missing)))
        return

    df = df.copy()
    for col in ["Standard Price", "Info Record Price", "PO Price"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["Report_Date"] = pd.to_datetime(df["Report_Date"], errors="coerce")
    df = df.dropna(subset=list(required))
    if df.empty:
        st.info("No hay datos válidos para perdidas.")
        return

    st.subheader("Dashboard de pérdida")
    st.caption("Analiza componentes donde Info Record Price > Standard Price.")

    min_d, max_d = df["Report_Date"].min().date(), df["Report_Date"].max().date()
    fc1, fc2, fc3 = st.columns([1, 1, 1.4])
    fc1.markdown("**Desde**")
    date_from = fc1.date_input(" ", value=min_d, min_value=min_d, max_value=max_d, key="loss_date_from", label_visibility="collapsed")
    fc2.markdown("**Hasta**")
    date_to   = fc2.date_input(" ", value=max_d, min_value=min_d, max_value=max_d, key="loss_date_to",   label_visibility="collapsed")
    if date_from > date_to:
        st.warning("La fecha de inicio no puede ser mayor a la fecha de fin.")
        return
    df = df[df["Report_Date"].between(pd.Timestamp(date_from), pd.Timestamp(date_to))].copy()
    sel = fc3.multiselect("Filtrar componentes", sorted(df["Component"].astype(str).unique()), key="loss_components")
    if sel:
        df = df[df["Component"].astype(str).isin(sel)]
    if df.empty:
        st.warning("Sin datos en el rango seleccionado.")
        return

    total = df["Component"].nunique()
    df["Loss Amount"]               = df["Info Record Price"] - df["Standard Price"]
    df["Loss Percent Vs Standard"]  = (df["Loss Amount"] / df["Standard Price"]) * 100
    df["PO vs Standard Diff"]       = df["PO Price"] - df["Standard Price"]
    df["PO vs Standard Percent"]    = (df["PO vs Standard Diff"] / df["Standard Price"]) * 100
    df = df[df["Loss Amount"] > 0].copy()
    if df.empty:
        st.success("No se encontraron pérdidas en el rango seleccionado.")
        return

    has_desc = "Component Description" in df.columns
    if has_desc:
        df["Component Description"] = df["Component Description"].fillna("")
    df = df.sort_values(["Component", "Report_Date"])
    df["Standard Price Previous"] = df.groupby("Component")["Standard Price"].shift(1)
    df["Variacion Porcentual"]    = ((df["Standard Price"] - df["Standard Price Previous"]) / df["Standard Price Previous"]) * 100

    aggs = {
        "Average_Loss":                ("Loss Amount", "mean"),
        "Average_Loss_Percent":        ("Loss Percent Vs Standard", "mean"),
        "Max_Loss":                    ("Loss Amount", "max"),
        "Average_Standard_Price":      ("Standard Price", "mean"),
        "Average_Info_Record_Price":   ("Info Record Price", "mean"),
        "Average_PO_Price":            ("PO Price", "mean"),
        "Average_PO_vs_Standard":      ("PO vs Standard Diff", "mean"),
        "Average_PO_vs_Standard_Percent": ("PO vs Standard Percent", "mean"),
        "Average_Variacion_Porcentual": ("Variacion Porcentual", "mean"),
        "Records":                     ("Component", "size"),
    }
    if has_desc:
        aggs["Component Description"] = ("Component Description", lambda v: next((x for x in v if x), ""))

    summary = (df.groupby("Component", as_index=False).agg(**aggs)
               .sort_values("Average_Loss", ascending=False).reset_index(drop=True))
    with_loss = summary["Component"].nunique()

    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Componentes totales", f"{total:,}")
    c2.metric("Con pérdida", f"{with_loss:,}", delta=f"{with_loss/total*100:.1f}% del total")
    c3.metric("Total pérdidas",     f"{df['Loss Amount'].sum():,.2f}")
    c4.metric("Pérdida promedio",   f"{df['Loss Amount'].mean():,.2f}")
    c5.metric("Pérdida máxima",     f"{df['Loss Amount'].max():,.2f}")

    top_n = st.slider("Top componentes", 1, max(5, min(50, len(summary))), min(15, len(summary)), key="loss_top_n")
    top   = summary.head(top_n)

    pie_df  = pd.DataFrame({"Categoria": ["Con pérdida", "Sin pérdida"], "Cantidad": [with_loss, max(total - with_loss, 0)]})
    pie_fig = px.pie(pie_df, names="Categoria", values="Cantidad", title="Distribución", hole=0.45,
                     color="Categoria", color_discrete_map={"Con pérdida": "#d62728", "Sin pérdida": "#2ca02c"})
    pie_fig.update_traces(textposition="inside", texttemplate="%{percent}",
                          hovertemplate="%{label}<br>%{value:,.0f}<br>%{percent}<extra></extra>")
    pie_fig.update_layout(height=max(420, top_n * 28), margin=dict(l=20, r=20, t=60, b=20))

    bar_fig = px.bar(top, x="Average_Loss", y="Component", orientation="h",
                     title="Componentes con mayor pérdida promedio",
                     hover_data=["Average_Loss_Percent", "Max_Loss", "Average_Standard_Price",
                                 "Average_PO_Price", "Average_PO_vs_Standard", "Average_Variacion_Porcentual", "Records"])
    bar_fig.update_traces(hovertemplate=(
        "Componente=%{y}<br>Perdida promedio=%{x:,.2f}<br>"
        "Perdida %=%{customdata[0]:,.2f}%<br>Perdida max=%{customdata[1]:,.2f}<br>"
        "Standard Price=%{customdata[2]:,.2f}<br>PO Price=%{customdata[3]:,.2f}<br>"
        "PO vs Standard=%{customdata[4]:,.2f}<br>Variacion %=%{customdata[5]:,.2f}%<br>"
        "Registros=%{customdata[6]:,.0f}<extra></extra>"
    ))
    bar_fig.update_layout(height=max(420, top_n * 28), yaxis={"categoryorder": "total ascending"},
                          margin=dict(l=20, r=20, t=60, b=20))

    cc1, cc2 = st.columns([1, 2])
    cc1.plotly_chart(pie_fig, use_container_width=True)
    cc2.plotly_chart(bar_fig, use_container_width=True)

    fmt = {
        "Average_Loss": "{:,.2f}", "Average_Loss_Percent": "{:,.2f}%", "Max_Loss": "{:,.2f}",
        "Average_Standard_Price": "{:,.2f}", "Average_Info_Record_Price": "{:,.2f}",
        "Average_PO_Price": "{:,.2f}", "Average_PO_vs_Standard": "{:,.2f}",
        "Average_PO_vs_Standard_Percent": "{:,.2f}%", "Average_Variacion_Porcentual": "{:,.2f}%",
        "Records": "{:,.0f}",
    }
    st.dataframe(top.reset_index(drop=True).style.format(fmt), use_container_width=True)

    sel_comp = st.selectbox("Detalle temporal por componente", summary["Component"].tolist(), key="loss_detail_component")
    det      = df[df["Component"] == sel_comp].copy()
    det_fig  = px.line(det, x="Report_Date", y=["Standard Price", "Info Record Price", "PO Price"],
                       markers=True, title=f"Evolución de precios — {sel_comp}")
    colors   = {"Standard Price": "#d62728", "Info Record Price": "#1f77b4", "PO Price": "#ff7f0e"}
    for trace in det_fig.data:
        trace.line.color   = colors.get(trace.name, "gray")
        trace.marker.color = colors.get(trace.name, "gray")
    det_fig.update_traces(hovertemplate="Fecha=%{x|%Y-%m-%d}<br>Valor=%{y:,.2f}<extra></extra>")
    det_fig.update_layout(hovermode="x unified")
    st.plotly_chart(det_fig, use_container_width=True)

    det_cols = (["Report_Date", "Component"] +
                (["Component Description"] if has_desc else []) +
                ["Standard Price", "Info Record Price", "PO Price",
                 "Loss Amount", "Loss Percent Vs Standard", "PO vs Standard Diff", "PO vs Standard Percent"])
    det_fmt  = {
        "Report_Date": lambda x: x.strftime("%Y-%m-%d") if pd.notnull(x) else "",
        "Standard Price": "{:,.2f}", "Info Record Price": "{:,.2f}", "PO Price": "{:,.2f}",
        "Loss Amount": "{:,.2f}", "Loss Percent Vs Standard": "{:,.2f}%",
        "PO vs Standard Diff": "{:,.2f}", "PO vs Standard Percent": "{:,.2f}%",
    }
    st.dataframe(det[det_cols].reset_index(drop=True).style.format(det_fmt), use_container_width=True)


def render_gain_dashboard(df):
    required = {"Report_Date", "Component", "Standard Price", "Info Record Price"}
    missing  = required.difference(df.columns)
    if missing:
        st.warning("Faltan columnas: " + ", ".join(sorted(missing)))
        return

    df = df.copy()
    for col in ["Standard Price", "Info Record Price"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["Report_Date"] = pd.to_datetime(df["Report_Date"], errors="coerce")
    df = df.dropna(subset=list(required))
    if df.empty:
        st.info("No hay datos válidos para ganancias.")
        return

    st.subheader("Dashboard de ganancia")
    st.caption("Analiza componentes donde Standard Price > Info Record Price.")

    min_d, max_d = df["Report_Date"].min().date(), df["Report_Date"].max().date()
    fc1, fc2, fc3 = st.columns([1, 1, 1.4])
    fc1.markdown("**Desde**")
    date_from = fc1.date_input(" ", value=min_d, min_value=min_d, max_value=max_d, key="gain_date_from", label_visibility="collapsed")
    fc2.markdown("**Hasta**")
    date_to   = fc2.date_input(" ", value=max_d, min_value=min_d, max_value=max_d, key="gain_date_to",   label_visibility="collapsed")
    if date_from > date_to:
        st.warning("La fecha de inicio no puede ser mayor a la fecha de fin.")
        return
    df = df[df["Report_Date"].between(pd.Timestamp(date_from), pd.Timestamp(date_to))].copy()
    sel = fc3.multiselect("Filtrar componentes", sorted(df["Component"].astype(str).unique()), key="gain_components")
    if sel:
        df = df[df["Component"].astype(str).isin(sel)]
    if df.empty:
        st.warning("Sin datos en el rango seleccionado.")
        return

    total      = df["Component"].nunique()
    status_df  = (
        df.assign(d=df["Standard Price"] - df["Info Record Price"])
        .groupby("Component", as_index=False)
        .agg(
            Has_Gain=("d", lambda v: (v > 0).any()),
            Has_Loss=("d", lambda v: (v < 0).any()),
            Is_Neutral=("d", lambda v: (v == 0).all()),
        )
    )
    neutral          = int(status_df["Is_Neutral"].sum())
    with_gain_total  = int(status_df["Has_Gain"].sum())
    loss_only        = int((~status_df["Has_Gain"] & status_df["Has_Loss"]).sum())

    df["Gain Amount"]               = df["Standard Price"] - df["Info Record Price"]
    df["Gain Percent Vs Info Record"] = (df["Gain Amount"] / df["Info Record Price"]) * 100
    df = df[df["Gain Amount"] > 0].copy()
    if df.empty:
        st.success("No se encontraron ganancias en el rango seleccionado.")
        return

    has_desc = "Component Description" in df.columns
    if has_desc:
        df["Component Description"] = df["Component Description"].fillna("")
    df = df.sort_values(["Component", "Report_Date"])
    df["Standard Price Previous"] = df.groupby("Component")["Standard Price"].shift(1)
    df["Variacion Porcentual"]    = ((df["Standard Price"] - df["Standard Price Previous"]) / df["Standard Price Previous"]) * 100

    aggs = {
        "Average_Gain":                 ("Gain Amount", "mean"),
        "Average_Gain_Percent":         ("Gain Percent Vs Info Record", "mean"),
        "Max_Gain":                     ("Gain Amount", "max"),
        "Average_Standard_Price":       ("Standard Price", "mean"),
        "Average_Info_Record_Price":    ("Info Record Price", "mean"),
        "Average_Variacion_Porcentual": ("Variacion Porcentual", "mean"),
        "Records":                      ("Component", "size"),
    }
    if has_desc:
        aggs["Component Description"] = ("Component Description", lambda v: next((x for x in v if x), ""))

    summary   = (df.groupby("Component", as_index=False).agg(**aggs)
                 .sort_values("Average_Gain", ascending=False).reset_index(drop=True))
    with_gain = summary["Component"].nunique()

    c1, c2, c3, c4, c5, c6 = st.columns(6)
    c1.metric("Componentes totales",  f"{total:,}")
    c2.metric("Con ganancia",         f"{with_gain_total:,}", delta=f"{with_gain_total/total*100:.1f}%")
    c3.metric("Neutrales",            f"{neutral:,}",         delta=f"{neutral/total*100:.1f}%")
    c4.metric("Total ganancias",      f"{df['Gain Amount'].sum():,.2f}")
    c5.metric("Ganancia promedio",    f"{df['Gain Amount'].mean():,.2f}")
    c6.metric("Ganancia máxima",      f"{df['Gain Amount'].max():,.2f}")

    top_n = st.slider("Top componentes", 1, max(5, min(50, len(summary))), min(15, len(summary)), key="gain_top_n")
    top   = summary.head(top_n)

    pie_df  = pd.DataFrame({
        "Categoria": ["Con ganancia", "Neutrales", "Solo pérdida"],
        "Cantidad":  [with_gain_total, neutral, loss_only],
    })
    pie_fig = px.pie(pie_df, names="Categoria", values="Cantidad", title="Distribución", hole=0.45,
                     color="Categoria",
                     color_discrete_map={"Con ganancia": "#2ca02c", "Neutrales": "#9fb3c8", "Solo pérdida": "#d62728"})
    pie_fig.update_traces(textposition="inside", texttemplate="%{percent}",
                          hovertemplate="%{label}<br>%{value:,.0f}<br>%{percent}<extra></extra>")
    pie_fig.update_layout(height=max(420, top_n * 28), margin=dict(l=20, r=20, t=60, b=20))

    bar_fig = px.bar(top, x="Average_Gain", y="Component", orientation="h",
                     title="Componentes con mayor ganancia promedio",
                     hover_data=["Average_Gain_Percent", "Max_Gain", "Average_Variacion_Porcentual", "Records"])
    bar_fig.update_traces(hovertemplate=(
        "Componente=%{y}<br>Ganancia promedio=%{x:,.2f}<br>"
        "Ganancia %=%{customdata[0]:,.2f}%<br>Ganancia max=%{customdata[1]:,.2f}<br>"
        "Variacion %=%{customdata[2]:,.2f}%<br>Registros=%{customdata[3]:,.0f}<extra></extra>"
    ))
    bar_fig.update_layout(height=max(420, top_n * 28), yaxis={"categoryorder": "total ascending"},
                          margin=dict(l=20, r=20, t=60, b=20))

    cc1, cc2 = st.columns([1, 2])
    cc1.plotly_chart(pie_fig, use_container_width=True)
    cc2.plotly_chart(bar_fig, use_container_width=True)

    fmt = {
        "Average_Gain": "{:,.2f}", "Average_Gain_Percent": "{:,.2f}%", "Max_Gain": "{:,.2f}",
        "Average_Standard_Price": "{:,.2f}", "Average_Info_Record_Price": "{:,.2f}",
        "Average_Variacion_Porcentual": "{:,.2f}%", "Records": "{:,.0f}",
    }
    st.dataframe(top.reset_index(drop=True).style.format(fmt), use_container_width=True)

    sel_comp = st.selectbox("Detalle temporal por componente", summary["Component"].tolist(), key="gain_detail_component")
    det      = df[df["Component"] == sel_comp].copy()
    det_fig  = px.line(det, x="Report_Date", y=["Standard Price", "Info Record Price"],
                       markers=True, title=f"Evolución de precios — {sel_comp}")
    up_color = "#2ca02c" if (det["Standard Price"] > det["Info Record Price"]).any() else "#1f77b4"
    colors   = {"Standard Price": up_color, "Info Record Price": "#1f77b4"}
    for trace in det_fig.data:
        trace.line.color   = colors.get(trace.name, "gray")
        trace.marker.color = colors.get(trace.name, "gray")
    det_fig.update_traces(hovertemplate="Fecha=%{x|%Y-%m-%d}<br>Valor=%{y:,.2f}<extra></extra>")
    det_fig.update_layout(hovermode="x unified")
    st.plotly_chart(det_fig, use_container_width=True)

    det_cols = (["Report_Date", "Component"] +
                (["Component Description"] if has_desc else []) +
                ["Standard Price", "Info Record Price", "Gain Amount", "Gain Percent Vs Info Record"])
    det_fmt  = {
        "Report_Date": lambda x: x.strftime("%Y-%m-%d") if pd.notnull(x) else "",
        "Standard Price": "{:,.2f}", "Info Record Price": "{:,.2f}",
        "Gain Amount": "{:,.2f}", "Gain Percent Vs Info Record": "{:,.2f}%",
    }
    st.dataframe(det[det_cols].reset_index(drop=True).style.format(det_fmt), use_container_width=True)


# ─── App entrypoint ───────────────────────────────────────────────────────────

st.markdown("## Control de Precios en Componentes")
st.caption("Consolida reportes semanales, limpia los datos y analiza variaciones de precio por componente.")

with st.sidebar:
    st.markdown("### Menú")
    st.caption("Sigue este flujo: Carga → Consolidado → Limpieza → Dashboards.")
    view_mode = st.radio(
        "Sección",
        ["Control de carga", "Consolidado", "Limpieza", "Dashboard de perdida", "Dashboard de ganancia"],
    )

if view_mode == "Control de carga":
    render_control_de_carga()

elif view_mode == "Consolidado":
    render_consolidado()

else:
    df_consolidated = st.session_state["consolidated_df"]
    if df_consolidated is None:
        st.info("Primero carga el consolidado desde la sección **Consolidado**.")
    else:
        df_clean, clean_metrics = clean_dataframe(df_consolidated)

        # Botón de exportar en sidebar (solo cuando hay datos limpios)
        with st.sidebar:
            st.markdown("---")
            ppt = build_powerpoint_export(df_clean)
            st.download_button(
                "Descargar PowerPoint", ppt,
                file_name=f"price_roll_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True,
            )
            st.markdown("---")
            toggle_label = "✕ Ocultar chat" if st.session_state["chat_open"] else "💬 Mostrar chat IA"
            if st.button(toggle_label, use_container_width=True, key="sidebar_toggle_chat"):
                st.session_state["chat_open"] = not st.session_state["chat_open"]
                st.rerun()

        # Reconstruir df_before_obs para la vista de limpieza
        n_before_obs = clean_metrics["rows_initial"] - clean_metrics["removed_material"] - clean_metrics["removed_demand"] - clean_metrics["removed_po"]
        df_before_obs = df_consolidated.copy()
        df_before_obs = df_before_obs[df_before_obs["Material Type"] != "ROH"]
        df_before_obs = df_before_obs[df_before_obs["Demand Comments"] != "Not found"]
        if "PO price comment" in df_before_obs.columns:
            df_before_obs = df_before_obs[df_before_obs["PO price comment"] != "PO price not found"]
        if "PO price comments" in df_before_obs.columns:
            df_before_obs = df_before_obs[df_before_obs["PO price comments"] != "PO price not found"]

        # Layout dinámico según visibilidad del chat
        if st.session_state["chat_open"]:
            cols = st.columns([3, 1.3])
        else:
            cols = [st.container()]

        with cols[0]:
            if view_mode == "Limpieza":
                render_limpieza(df_clean, df_before_obs, clean_metrics)
            elif view_mode == "Dashboard de perdida":
                render_loss_dashboard(df_clean)
            elif view_mode == "Dashboard de ganancia":
                render_gain_dashboard(df_clean)

        if st.session_state["chat_open"]:
            extra_tables = {
                "Tabla intermedia de limpieza (sin ROH/Demand/PO, con OBS/IN aún)": df_before_obs,
                "Metadata de archivos procesados": load_metadata(),
                "Diagnóstico de archivos en carpeta": get_diagnostics_report(),
            }
            with cols[1]:
                render_chat_panel(df_clean, df_consolidated, extra_tables)
